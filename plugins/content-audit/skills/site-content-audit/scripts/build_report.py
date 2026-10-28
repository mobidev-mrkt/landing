#!/usr/bin/env python3
"""Step 8. Build the Excel deliverable.

Usage: python3 -I build_report.py --work WORK --out report.xlsx
Inputs (WORK): findings.json, auto.json, txt/_meta.json, schema_pages.json, schema_site.json,
  optional spell_confirmed.json {"wrong": "right"} and facts.json [[fact, version A, version B, action], ...]
  (facts.json is curated by hand from cross_facts.json — contradictions between pages).
Sheets: Зведення · Знахідки (сайт) · Знахідки (блог) · Факти компанії · Title і meta · Schema — весь сайт · Schema — по сторінках · Schema — шаблони
Every findings row has an exact quote for Ctrl+F, a fix, and a Status dropdown.
"""
import argparse, os, re, json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

CAT = {'typo': 'Описка', 'grammar': 'Граматика', 'placeholder': 'Плейсхолдер / незавершене', 'broken_text': 'Зламаний текст',
       'inconsistency': 'Нестиковка', 'outdated': 'Застаріле', 'meta': 'Title / meta description'}
PRI = {'high': '1 Високий', 'medium': '2 Середній', 'low': '3 Низький'}
BLOG_SEGMENTS = ('blog', 'articles', 'insights', 'news', 'posts', 'resources', 'journal', 'stories', 'trends-and-insights')
SECT = {'services': 'Сервіси', 'service': 'Сервіси', 'solutions': 'Сервіси', 'case-studies': 'Кейси', 'cases': 'Кейси',
        'portfolio': 'Кейси', 'success-stories': 'Кейси', 'our-team': 'Команда', 'team': 'Команда', 'authors': 'Команда',
        'webinar': 'Вебінари', 'webinars': 'Вебінари', 'events': 'Події', 'category': 'Категорії', 'industries': 'Галузі'}
ORDER = ['Шаблон (багато сторінок)', 'Головна', 'Сервіси', 'Кейси', 'Галузі', 'Лендінги / інше']
TAIL = ['Команда', 'Блог']
HERE = os.path.dirname(os.path.abspath(__file__))

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--work', required=True); ap.add_argument('--out', required=True)
    a = ap.parse_args(); W = a.work; site = json.load(open(f'{W}/site.json'))['site'] + '/'
    load = lambda f, d=None: json.load(open(f'{W}/{f}')) if os.path.exists(f'{W}/{f}') else d
    url = lambda p: p if p.startswith('http') else site + ('' if p == 'home' else p.replace('__', '/'))
    def sect(u):
        parts = u.replace(site, '').strip('/').split('/')
        if parts == ['']: return 'Головна'
        if len(parts) == 1: return SECT.get(parts[0], 'Лендінги / інше')
        if parts[0] in BLOG_SEGMENTS: return 'Блог'
        return SECT.get(parts[0], parts[0].replace('-', ' ').capitalize())
    rows = []
    for r in load('findings.json', []):
        if r.get('category') not in CAT or r.get('severity') not in PRI: continue
        note, sec = r.get('note', ''), sect(r['url'])
        if r.get('template_pages'):
            tp = r['template_pages']; sec = 'Шаблон (багато сторінок)'
            note = f"Повторюється на {len(tp)} сторінках — елемент шаблону, виправити один раз. Напр.: " + ', '.join(u.replace(site, '/') for u in tp[:3]) + '. ' + note
        rows.append([PRI[r['severity']], sec, r['url'], CAT[r['category']], r['quote'], r.get('fix', ''), note, 'Вичитка'])
    have = collections.defaultdict(str)
    for r in rows: have[r[2]] += ' ' + r[4]
    A = load('auto.json', {})
    def auto(pri, p, cat, quote, fix, note=''):
        u = url(p)
        if quote[:25] not in have[u]: rows.append([pri, sect(u), u, cat, quote, fix, note, 'Автоперевірка']); have[u] += ' ' + quote
    for p, t in A.get('cyrillic_in_latin', []):
        w, ctx = t.split(' :: ', 1); auto(PRI['medium'], p, 'Кирилиця в англ. слові', ctx, f'Перенабрати «{w}» латиницею', 'Слово з кириличною літерою не знаходиться пошуком')
    spell = load('spell_confirmed.json', {})
    for w, right, p, ctx in load('spell_candidates.json', []):
        if w in spell: auto(PRI['medium'], p, 'Описка', ctx, f'{w} → {spell[w]}')
    for p, t in A.get('alt_number_missing', []): auto(PRI['medium'], p, 'Alt без цифри', t, 'Повернути цифри в alt', 'Баг шаблону: alt генерується з назви файлу')
    for p, t in A.get('mojibake', []): auto(PRI['high'], p, 'Зламаний текст', t, 'Виправити кодування')
    order = {v: i for i, v in enumerate(PRI.values())}
    rank = lambda sec: ORDER.index(sec) if sec in ORDER else (100 + TAIL.index(sec) if sec in TAIL else 50)
    rows.sort(key=lambda r: (order[r[0]], rank(r[1]), r[1], r[2]))
    wb = Workbook(); su = wb.active; su.title = 'Зведення'
    def sheet(title, hdr, data, widths, link_col=None, status=False):
        ws = wb.create_sheet(title); ws.append((['#'] if status else []) + hdr + (['Статус'] if status else []))
        for i, r in enumerate(data, 1):
            ws.append(([i] if status else []) + list(r) + (['Не почато'] if status else []))
            if link_col:
                c = ws.cell(i + 1, link_col + (1 if status else 0)); c.hyperlink = c.value; c.font = Font(color='1155CC', underline='single')
        for i, w in enumerate(([5] if status else []) + widths + ([12] if status else []), 1): ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2937')
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        if status and data:
            dv = DataValidation(type='list', formula1='"Не почато,В роботі,Виправлено,Не баг"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f'{get_column_letter(ws.max_column)}2:{get_column_letter(ws.max_column)}{len(data) + 1}')
        return ws
    hdr = ['Пріоритет', 'Розділ', 'Сторінка', 'Тип', 'Що не так (точна цитата для Ctrl+F)', 'Як виправити', 'Коментар', 'Як знайдено']
    site_rows_, blog_rows = [r for r in rows if r[1] != 'Блог'], [r for r in rows if r[1] == 'Блог']
    for title, data in (('Знахідки (сайт)', site_rows_), ('Знахідки (блог)', blog_rows)):
        ws = sheet(title, hdr, data, [12, 14, 48, 22, 60, 45, 40, 14], link_col=3, status=True)
        for row in ws.iter_rows(min_row=2): row[1].fill = PatternFill('solid', fgColor={'1': 'FBEBEB', '2': 'F9EEDA', '3': 'F0EFEC'}[row[1].value[0]])
    sheet('Факти компанії', ['Факт', 'Версія А', 'Версія Б (і далі)', 'Що зробити'], load('facts.json', []), [22, 55, 75, 45])
    meta = load('txt/_meta.json', {}); seo = []
    for k in ('meta_missing_desc', 'meta_missing_title', 'meta_title_long', 'meta_desc_long', 'meta_dup_title', 'meta_dup_desc', 'h1_count'):
        for p, t in A.get(k, []): seo.append([url(p), sect(url(p)), k, t])
    sheet('Title і meta', ['Сторінка', 'Розділ', 'Проблема', 'Деталі'], seo, [70, 16, 22, 80], link_col=1)
    sr = load('schema_site.json', []); sp = load('schema_pages.json', [])
    sheet('Schema — весь сайт', ['Пріоритет', 'Де', 'Проблема', 'Що зробити', 'Навіщо (AI-пошук / Google)'], sr, [12, 28, 60, 70, 60])
    sp = sorted(([pr, sect(url(p)), url(p), prob, fix, why] for p, pr, prob, fix, why in sp), key=lambda r: (r[0], r[1], r[2]))
    sheet('Schema — по сторінках', ['Пріоритет', 'Розділ', 'Сторінка', 'Проблема', 'Що зробити', 'Навіщо'], sp, [12, 14, 55, 50, 60, 55], link_col=3, status=True)
    tpl = json.load(open(os.path.join(HERE, '..', 'references', 'schema_templates.json')))
    sheet('Schema — шаблони', ['Тип сторінки', 'JSON-LD (вставити в <script type="application/ld+json">, замінити <...>)'],
          [[k, json.dumps(v, ensure_ascii=False, indent=2)] for k, v in tpl.items()], [28, 120])
    su.append([f'Аудит контенту {site}']); su['A1'].font = Font(bold=True, size=14); su.append([])
    su.append(['Сторінок у sitemap', len(meta)]); su.append([])
    su.append(['Тип', 'Високий', 'Середній', 'Низький', 'Разом'])
    cc = collections.Counter((r[3], r[0][0]) for r in rows)
    for c in sorted({r[3] for r in rows}, key=lambda c: -sum(cc[(c, x)] for x in '123')):
        x = [cc[(c, k)] for k in '123']; su.append([c, *x, sum(x)])
    su.append(['Разом', *[sum(1 for r in rows if r[0][0] == k) for k in '123'], len(rows)])
    su.append(['  з них «Знахідки (сайт)»', *[sum(1 for r in site_rows_ if r[0][0] == k) for k in '123'], len(site_rows_)])
    su.append(['  з них «Знахідки (блог)»', *[sum(1 for r in blog_rows if r[0][0] == k) for k in '123'], len(blog_rows)])
    su.append([]); su.append(['Title і meta', len(seo)]); su.append(['Schema — по сторінках', len(sp)]); su.append(['Schema — весь сайт', len(sr)])
    su.column_dimensions['A'].width = 40
    wb.save(a.out); print('saved', a.out, len(rows), 'findings')

if __name__ == '__main__':
    main()
